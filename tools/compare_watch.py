#!/usr/bin/env python3
"""Confronter la veille de l'agent à la veille presse du cabinet.

    python3 tools/compare_watch.py              # la comparaison, hors ligne
    python3 tools/compare_watch.py --probe      # sonde en plus les domaines

Le cabinet reçoit une veille presse (agrégateur de type Factiva) sur les mêmes
sujets que l'agent. Les deux listes ne se recouvrent presque pas, et la question
posée est : pourquoi l'agent n'a-t-il pas ces articles.

Trois causes possibles, qu'il faut séparer avant de conclure quoi que ce soit :

  fenêtre     l'agent n'a pas tourné sur la période. Rien à corriger dans le
              scraping ; c'est une question d'exploitation.
  registre    le domaine n'est pas dans tblSources. L'agent ne l'a pas raté :
              il ne l'a jamais regardé.
  collecte    le domaine est surveillé mais la page n'a pas pu être lue.

Et une quatrième, qui va dans l'autre sens : une bonne part de la liste du
cabinet n'a rien à voir avec NIS 2. Un agrégateur qui filtre sur la chaîne
« Network and Information Security » ramène des appels d'offres chinois, et
« DSP » ramène du traitement du signal audio. Compter ces lignes comme des
manques de l'agent fausserait la mesure dans l'autre sens.

La liste du cabinet est saisie ici telle que reçue, avec pour chaque ligne le
motif de son classement, pour que le jugement soit relisible et discutable.
"""

import argparse
import json
import sys
import urllib.parse as up
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
ITEMS = ROOT / "data" / "watch-items.json"
WORKBOOK = ROOT / "ressources" / "agent_veille_NIS2.xlsx"

# ---------------------------------------------------------------------------
# La veille du cabinet, reçue le 7 septembre 2026.
#
#   sujet   la ligne parle de la transposition, d'une autorité, d'une
#           obligation : c'est ce que l'outil suit.
#   marge   cyber ou conformité, mais ni transposition ni autorité - un produit,
#           un événement, une levée de fonds. Utile à un commercial, pas au
#           classeur comparatif.
#   bruit   aucun rapport : l'agrégateur a filtré sur une chaîne de caractères.
# ---------------------------------------------------------------------------
FIRM = [
    # (date, source, domaine, classement, titre court, motif du classement)
    ("2026-09-07", "GoodTech.info", "goodtech.info", "marge",
     "SecNumCloud : Numspot décroche la qualification ANSSI",
     "qualification SecNumCloud, pas la transposition NIS 2"),
    ("2026-09-07", "France 3", "france3-regions.francetvinfo.fr", "sujet",
     "Le gouvernement n'a toujours pas fait voter NIS 2",
     "état de la transposition française, et calendrier parlementaire"),
    ("2026-09-04", "MediaDreams.fr", "mediadreams.fr", "marge",
     "Constellation Experience Day 2026", "événement éditeur"),
    ("2026-09-04", "Generation-NT", "generation-nt.com", "marge",
     "Le défi numérique européen, de la régulation à la souveraineté",
     "souveraineté numérique, NIS2 cité en passant"),
    ("2026-09-04", "Techniques de l'Ingénieur", "techniques-ingenieur.fr", "sujet",
     "NIS2, RGPD et normes ISO, les fondamentaux de l'audit",
     "obligations d'audit et responsabilité des dirigeants"),
    ("2026-09-03", "Zonebourse", "zonebourse.com", "sujet",
     "KnowBe4 accrédité par l'ACN italienne",
     "schéma d'accréditation d'une autorité nationale"),
    ("2026-09-03", "VOI.id", "voi.id", "bruit",
     "Culture du droit et de l'éthique, Karaton Surakarta",
     "Indonésie, justice locale - filtré sur « directives »"),
    ("2026-09-02", "La Gazette France", "lagazettefrance.fr", "marge",
     "Normandie : les entreprises face aux cyberattaques", "webinaire régional"),
    ("2026-09-02", "L'Informaticien.com", "linformaticien.com", "marge",
     "Cohesity ouvre une clean room dans HCLTech VaultNXT", "annonce produit"),
    ("2026-09-01", "France Info", "francetvinfo.fr", "sujet",
     "Comment les hackers s'introduisent dans les systèmes de l'État",
     "enquête : la France tarde à transposer NIS 2"),
    ("2026-09-01", "Archimag", "archimag.com", "marge",
     "Documents, données, IA : bâtir une chaîne de preuve conforme",
     "conformité documentaire, NIS 2 dans une liste"),
    ("2026-09-01", "Fédération Hospitalière de France", "fhf.fr", "marge",
     "Cybersécurité et services de santé", "appel européen DIGITAL-ECCC"),
    ("2026-08-31", "VOI.id", "voi.id", "bruit",
     "Transactions roupie / dollar singapourien", "aucun rapport"),

    ("2026-09-04", "Professional Security", "professionalsecurity.co.uk", "marge",
     "Optimal Crisis Leaders", "lancement d'offre, NIS2 et CER cités"),
    ("2026-09-04", "Polish Press Agency", "pap.pl", "marge",
     "Vercom IT entre sur le marché de la cybersécurité", "stratégie d'éditeur"),
    ("2026-09-03", "Gov.uk Find a Tender", "find-tender.service.gov.uk", "bruit",
     "Quality Delivery Partner Services", "appel d'offres défense - filtré sur « DSP »"),
    ("2026-09-03", "Broadcast Bridge", "thebroadcastbridge.com", "bruit",
     "Nugen Audio lance ses abonnements", "traitement du signal audio - « DSP »"),
    ("2026-09-03", "ENP Newswire", "enpnewswire.com", "marge",
     "Nouvel appel ECCC, 96 M€", "financement européen, pas une transposition"),
    ("2026-09-03", "SustainabilityMag", "sustainabilitymag.com", "bruit",
     "Carbmee et la durabilité", "aucun rapport"),
    ("2026-09-03", "MarketLine", "marketline.com", "bruit",
     "EIS Holdings acquiert Environmental Construction", "M&A - filtré sur « EIS »"),
    ("2026-09-02", "MarketLine", "marketline.com", "bruit",
     "Viatel acquiert EDNX", "M&A"),
    ("2026-09-02", "ENP Newswire", "enpnewswire.com", "marge",
     "Cohesity clean room pour HCLTech", "annonce produit, doublon FR"),
    ("2026-09-02", "EMIS Insights", "emis.com", "sujet",
     "Nouveau paysage cyber en Pologne - mise en oeuvre de NIS2",
     "rapport sur la transposition polonaise"),
    ("2026-09-02", "BusinessPlus.ie", "businessplus.ie", "bruit",
     "Viatel rachète EDNX", "M&A, doublon"),
    ("2026-09-02", "MarketLine", "marketline.com", "bruit",
     "Xorlab lève 5,8 M$", "levée de fonds"),
    ("2026-09-01", "Aberdeen & Grampian Chamber", "agcc.co.uk", "bruit",
     "OES Group repasse en mains privées", "rachat, secteur pétrolier"),
    ("2026-09-01", "Nordic Daily", "nordicdaily.com", "sujet",
     "Rapport final d'évaluation de la loi cyber finlandaise",
     "évaluation officielle de la transposition finlandaise"),
    ("2026-09-01", "Help Net Security", "helpnetsecurity.com", "sujet",
     "Conformité NIS2 : IAM et contrôle d'accès avant l'audit 2026",
     "obligations et échéance d'audit d'octobre"),
    ("2026-09-01", "Tech.eu", "tech.eu", "bruit",
     "xorlab lève 5 M€", "levée de fonds"),
    ("2026-08-31", "LVM.fi", "lvm.fi", "sujet",
     "La loi cyber finlandaise jugée globalement réussie",
     "source officielle : ministère finlandais des Transports et Communications"),
    ("2026-08-31", "PortSEurope", "portseurope.com", "sujet",
     "Les ports néerlandais soumis aux nouvelles lois cyber",
     "entrée en vigueur des Cbw et Wwke le 15 août - transposition NL de NIS2 et CER"),
    ("2026-08-31", "EuropaWire", "europawire.eu", "bruit",
     "Chercheurs de Groningue sur Anthropic Mythos", "aucun rapport"),
    ("2026-08-31", "GistMania", "gistmania.com", "bruit",
     "Police démantèle un réseau au Nigeria", "filtré sur « DSP »"),
    ("2026-08-31", "Aberdeen Business News", "aberdeenbusinessnews.co.uk", "bruit",
     "OES Group, rachat", "doublon"),
    ("2026-08-31", "BNS", "bns.ee", "sujet",
     "Subventions cyber pour les entités nouvellement assujetties (Estonie)",
     "dispositif d'accompagnement lié à l'assujettissement NIS 2"),
    ("2026-08-26", "Global Data Point", "globaldatapoint.com", "marge",
     "iluminr transforme les tests de scénarios", "annonce produit"),

    ("2026-09-05", "Tendernews.com", "tendernews.com", "bruit",
     "Appel d'offres sécurité réseau, Xinjiang", "marché public chinois"),
    ("2026-09-05", "Tenders Monitor", "tendersmonitor.com", "bruit",
     "Appel d'offres sécurité réseau", "marché public chinois"),
    ("2026-09-04", "Tendernews.com", "tendernews.com", "bruit",
     "Équipements réseau, Kunming", "marché public chinois"),
    ("2026-09-04", "Tendernews.com", "tendernews.com", "bruit",
     "Centre de transport du Xinjiang", "marché public chinois"),
    ("2026-09-04", "EUROPE SAYS", "europesays.com", "marge",
     "Le corpus numérique de l'UE à un tournant - CEPS",
     "débat sur la simplification réglementaire"),
    ("2026-09-03", "Electronics Daily", "electronicsdaily.com", "bruit",
     "Reconstruction de veines palmaires, Shandong", "recherche académique"),
    ("2026-09-03", "GeekWire", "geekwire.com", "bruit",
     "Qualtrics supprime 117 postes", "filtré sur un intitulé de poste"),
    ("2026-09-02", "Tendernews.com", "tendernews.com", "bruit",
     "Support technique, Qinghai", "marché public chinois"),
    ("2026-09-02", "Tendernews.com", "tendernews.com", "bruit",
     "Plateforme de formation, Shaanxi", "marché public chinois"),
    ("2026-09-01", "Arthur Cox", "arthurcox.com", "sujet",
     "Guidance du NCSC irlandais sur la gouvernance cyber des entités NIS2",
     "lecture d'une publication d'autorité nationale"),
    ("2026-09-01", "Tendernews.com", "tendernews.com", "bruit",
     "Maintenance, Hunan", "marché public chinois"),
    ("2026-09-01", "Tendernews.com", "tendernews.com", "bruit",
     "Rectificatif, Shaanxi", "marché public chinois"),
    ("2026-08-26", "Tenders Monitor", "tendersmonitor.com", "bruit",
     "Équipements, Corée", "marché public coréen"),
    ("2026-08-25", "MarketLine", "marketline.com", "bruit",
     "Shanghai Yufeng Tunan lève des fonds", "levée de fonds"),
    ("2026-08-20", "MarketLine", "marketline.com", "bruit",
     "Acquisition Keneng Tengda", "M&A"),
]

# Les autorités que ces articles citent. Si l'agent surveille l'autorité, il
# pouvait tenir le fait sans passer par la presse - c'est le test qui compte.
BEHIND = {
    "2026-08-31|portseurope.com": ("NL", "ncsc.nl", "entrée en vigueur Cbw + Wwke"),
    "2026-08-31|lvm.fi": ("FI", "lvm.fi", "rapport d'évaluation du ministère"),
    "2026-09-01|nordicdaily.com": ("FI", "lvm.fi", "même rapport, repris"),
    "2026-09-01|arthurcox.com": ("IE", "ncsc.gov.ie", "guidance gouvernance du NCSC-IE"),
    "2026-09-03|zonebourse.com": ("IT", "acn.gov.it", "accréditation ACN"),
}


def registry_hosts():
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("pip install openpyxl")
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb["Sources"]
    head = [c.value for c in ws[1]]
    col = {str(h).strip(): i for i, h in enumerate(head) if h}
    hosts = set()
    for r in range(2, ws.max_row + 1):
        url = ws.cell(row=r, column=col["URL / Endpoint"] + 1).value
        if url:
            hosts.add(up.urlparse(str(url)).netloc.lower().removeprefix("www."))
    return hosts


def collected_hosts():
    items = json.loads(ITEMS.read_text(encoding="utf-8"))["items"]
    hosts = Counter()
    for i in items:
        u = (i.get("source") or {}).get("url") or ""
        hosts[up.urlparse(u).netloc.lower().removeprefix("www.")] += 1
    dates = sorted(i["detected"] for i in items)
    return hosts, dates[0], dates[-1]


# Un site peut très bien publier un flux sans l'annoncer dans son <head>.
# Ne tester que le <link> déclaré donnait "aucun flux" sur des domaines qui en
# servent un à l'adresse la plus banale : la conclusion aurait été fausse.
FEED_PATHS = ("/feed", "/feed/", "/rss", "/rss.xml", "/feed.xml", "/atom.xml",
              "/index.xml", "/en/rss", "/actualites/rss", "/titres.rss")


def probe(host):
    """Le domaine répond-il, et publie-t-il un flux ? Mesuré, pas supposé."""
    import requests
    from bs4 import BeautifulSoup
    ua = {"User-Agent": "Mozilla/5.0 (compatible; RegWatch; +internal)"}

    def parses(url):
        """Un flux n'est un flux que s'il s'analyse et porte des entrées."""
        try:
            import feedparser
            rr = requests.get(url, timeout=12, headers=ua)
            if rr.status_code != 200:
                return None
            d = feedparser.parse(rr.content)
            return len(d.entries) if d.entries else None
        except Exception:
            return None

    try:
        r = requests.get("https://" + host, timeout=12, headers=ua, allow_redirects=True)
    except Exception as e:
        return {"status": type(e).__name__, "feed": None, "entries": None}
    out = {"status": r.status_code, "feed": None, "entries": None}
    if r.status_code != 200:
        return out

    tries = []
    soup = BeautifulSoup(r.text, "html.parser")
    for link in soup.find_all("link"):
        ty = (link.get("type") or "").lower()
        if "rss" in ty or "atom" in ty:
            if link.get("href"):
                tries.append(up.urljoin(r.url, link["href"]))
    tries += [up.urljoin(r.url, p) for p in FEED_PATHS]

    for url in tries:
        n = parses(url)
        if n:
            out["feed"], out["entries"] = url, n
            break
    return out


# Ce que le sondage a mesuré le 7 septembre 2026. Ecrit ici plutot que resonde
# a chaque appel : ces lignes partent dans un classeur, elles doivent etre
# stables et relisibles, pas dependre de l'humeur d'un serveur.
MEASURED_FEEDS = {
    "helpnetsecurity.com": "https://www.helpnetsecurity.com/feed/",
    "portseurope.com": "https://www.portseurope.com/feed",
    "arthurcox.com": "https://www.arthurcox.com/feed",
    "bns.ee": "https://www.bns.ee/rss.xml",
    "francetvinfo.fr": "https://www.franceinfo.fr/titres.rss",
    "france3-regions.francetvinfo.fr": "https://france3-regions.franceinfo.fr/en/rss",
    "emis.com": "https://isimarkets.com/feed/",
}
# Sondes sans flux : a surveiller en page, ou a laisser tomber.
NO_FEED = {
    "lvm.fi": "ministère finlandais - publie les évaluations de la loi cyber, sans flux",
    "nordicdaily.com": "reprise de presse, sans flux",
    "techniques-ingenieur.fr": "presse technique, sans flux exploitable",
    "zonebourse.com": "403 : pare-feu applicatif, inatteignable sans contournement",
}

OUT_CSV = ROOT / "data" / "tblSources-presse.csv"


def write_rows(missing):
    """Les lignes a coller dans tblSources, au format du classeur de l'agent."""
    import csv
    cols = ["Source", "Type", "URL / Endpoint", "Actif", "Pays / zone",
            "Fiabilité", "Priorité", "Note"]
    named = {
        "helpnetsecurity.com": ("Help Net Security", "EU"),
        "portseurope.com": ("PortSEurope", "EU"),
        "arthurcox.com": ("Arthur Cox - veille juridique", "IE"),
        "bns.ee": ("Baltic News Service", "EE"),
        "francetvinfo.fr": ("France Info", "FR"),
        "france3-regions.francetvinfo.fr": ("France 3 Régions", "FR"),
        "emis.com": ("EMIS Insights", "EU"),
    }
    rows = []
    for host, feed in MEASURED_FEEDS.items():
        name, iso = named[host]
        rows.append({
            "Source": name, "Type": "RSS", "URL / Endpoint": feed, "Actif": "Oui",
            "Pays / zone": iso, "Fiabilité": "Non officielle - à vérifier",
            "Priorité": "2",
            # Le degre de confiance n'est pas une opinion : ces domaines ne sont
            # pas des autorites, tout ce qu'ils rapportent demande une source
            # officielle avant publication.
            "Note": "Presse ou cabinet - alerte précoce, à confirmer sur la source officielle. "
                    "Flux vérifié le 7 septembre 2026.",
        })
    for host, why in NO_FEED.items():
        rows.append({
            "Source": host, "Type": "Page web", "URL / Endpoint": "https://" + host,
            "Actif": "Non", "Pays / zone": "", "Fiabilité": "Non officielle - à vérifier",
            "Priorité": "3", "Note": why + " - inactive tant que la collecte n'est pas traitée.",
        })
    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, delimiter=";")
        w.writeheader()
        w.writerows(rows)
    print("\n%d lignes -> %s" % (len(rows), OUT_CSV.relative_to(ROOT)))
    print("  %d avec flux vérifié (actives), %d sans flux (inactives, documentées)"
          % (len(MEASURED_FEEDS), len(NO_FEED)))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--probe", action="store_true", help="sonder les domaines « sujet »")
    ap.add_argument("--rows", action="store_true",
                    help="écrire les lignes tblSources des domaines à ajouter")
    args = ap.parse_args()

    reg = registry_hosts()
    got, first, last = collected_hosts()

    kinds = Counter(k for _d, _s, _h, k, _t, _m in FIRM)
    print("VEILLE DU CABINET - %d lignes reçues" % len(FIRM))
    for k in ("sujet", "marge", "bruit"):
        print("  %-6s %2d  (%d %%)" % (k, kinds[k], round(100 * kinds[k] / len(FIRM))))
    print("\n  Le filtre du cabinet est textuel : « Network and Information Security »")
    print("  ramène des marchés publics chinois, « DSP » du traitement audio.")

    print("\nAGENT - dernière collecte")
    print("  détections du %s au %s" % (first, last))
    late = [d for d, *_ in FIRM if d > last]
    print("  %d des %d lignes du cabinet sont postérieures au dernier run"
          % (len(late), len(FIRM)))

    subjects = [x for x in FIRM if x[3] == "sujet"]
    print("\nLES %d LIGNES QUI RELÈVENT DU SUJET" % len(subjects))
    print("  %-12s %-30s %-9s %s" % ("date", "domaine", "registre", "titre"))
    for d, _s, h, _k, title, _m in subjects:
        print("  %-12s %-30s %-9s %s"
              % (d, h[:30], "oui" if h in reg else "NON", title[:44]))

    missing = sorted({h for _d, _s, h, k, _t, _m in FIRM if k == "sujet" and h not in reg})
    print("\n  %d domaine(s) « sujet » absent(s) du registre : %s"
          % (len(missing), ", ".join(missing)))

    print("\nCE QUE L'AGENT POUVAIT TENIR SANS LA PRESSE")
    for key, val in BEHIND.items():
        if not val:
            continue
        date, host = key.split("|")
        iso, auth, what = val
        watched = auth in reg
        print("  %-4s %-18s %-9s %s" % (iso, auth, "surveillé" if watched else "ABSENT", what))

    if args.probe:
        print("\nSONDAGE DES DOMAINES « SUJET » ABSENTS DU REGISTRE")
        for h in missing:
            r = probe(h)
            print("  %-31s %-9s %s" % (h, r["status"],
                  ("%s (%d entrées)" % (r["feed"], r["entries"])) if r["feed"] else "aucun flux"))
    else:
        print("\n--probe pour sonder les domaines absents (réseau).")

    if args.rows:
        write_rows(missing)
    return 0


if __name__ == "__main__":
    sys.exit(main())
