#!/usr/bin/env python3
"""Extract the workbook's "Mapping cyber themes" sheet into the country records.

    python3 tools/cyber_themes.py "ressources/CYBER WATCH5_Technical inventory.xlsx"

The sheet answers a question none of the other tabs do: not "how many cyber
requirements does this country impose", which the KPI table already gives, but
"WHICH themes does its framework actually cover". Two countries with 150
requirements each can cover completely different ground, and that is what a
consultant needs before telling a client where the gaps are.

Shape of the source, from mapping it:
  row 9    family name, written once above its block of columns
  row 10   theme name, one per column
  row 11+  one country per row, Yes / No per theme
  the last column of each family is a TOTAL: the family is covered at all

Two limits that travel with the data, and are surfaced rather than hidden:
  - 9 countries out of 29. The sheet was never completed.
  - last updated 30/07/2025, where the rest of the workbook is June-July 2026.
    An absent theme may simply not have been reviewed since.

Output: data/cyber-themes.json, src/reg/nis2/data_themes.js
"""

import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit("pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
SHEET = "Mapping cyber themes"
GROUP_ROW, HEAD_ROW, FIRST_ROW = 9, 10, 11
COUNTRY_COL = 3
TOTAL = "total"

# The sheet spells a few countries its own way.
COUNTRY_FIX = {"Czech Republic": "Czechia", "UK": "United Kingdom"}


def clean(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    return re.sub(r"\s+", " ", str(value)).strip()


def yes(value):
    return clean(value).lower() in ("yes", "oui", "y", "true", "1")


def sheet_date(ws):
    """The sheet states its own age in A2; it is older than the workbook."""
    for row in range(1, GROUP_ROW):
        for cell in ws[row]:
            m = re.search(r"(\d{2})/(\d{2})/(\d{4})", clean(cell.value))
            if m:
                return "%s-%s-%s" % (m.group(3), m.group(2), m.group(1))
    return ""


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    path = Path(sys.argv[1]).expanduser()
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit('feuille "%s" absente (présentes : %s)'
                         % (SHEET, ", ".join(wb.sheetnames)))
    ws = wb[SHEET]

    # A family name is written once, above the first of its columns.
    family_of, current = {}, None
    for col in range(1, ws.max_column + 1):
        label = clean(ws.cell(row=GROUP_ROW, column=col).value)
        if label:
            current = label
        family_of[col] = current

    themes = []          # (col, family, theme) - TOTAL columns kept apart
    totals = {}          # family -> col
    for col in range(COUNTRY_COL + 3, ws.max_column + 1):
        head = clean(ws.cell(row=HEAD_ROW, column=col).value)
        family = family_of.get(col)
        if not head or not family:
            continue
        if head.lower() == TOTAL:
            totals[family] = col
        else:
            themes.append((col, family, head))

    families = []
    for _col, family, theme in themes:
        entry = next((f for f in families if f["family"] == family), None)
        if not entry:
            entry = {"family": family, "themes": []}
            families.append(entry)
        entry["themes"].append(theme)

    # The country records already carry the ISO code and the English name.
    known = {c["name"]: c["iso"] for c in json.loads(
        (ROOT / "data" / "countries.json").read_text(encoding="utf-8"))["countries"]}

    countries, unknown = {}, []
    for row in range(FIRST_ROW, ws.max_row + 1):
        name = clean(ws.cell(row=row, column=COUNTRY_COL).value)
        if not name:
            continue
        name = COUNTRY_FIX.get(name, name)
        iso = known.get(name)
        if not iso:
            unknown.append(name)
            continue
        covered, blocks = [], []
        for entry in families:
            hit = [t for col, fam, t in themes
                   if fam == entry["family"] and yes(ws.cell(row=row, column=col).value)]
            covered.extend(hit)
            total_col = totals.get(entry["family"])
            blocks.append({
                "family": entry["family"],
                "covered": len(hit),
                "of": len(entry["themes"]),
                "themes": hit,
                # The sheet's own verdict for the family, kept because it is
                # what the consultant filled in - not recomputed from the row.
                "familyCovered": yes(ws.cell(row=row, column=total_col).value) if total_col else None,
            })
        countries[iso] = {
            "country": name,
            "themesCovered": len(covered),
            "themesTotal": len(themes),
            "families": blocks,
        }

    payload = {
        "workbook": path.name,
        "sheet": SHEET,
        "lastUpdate": sheet_date(ws),
        "families": families,
        "themeCount": len(themes),
        "countries": countries,
    }
    (ROOT / "data" / "cyber-themes.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    out = ROOT / "src" / "reg" / "nis2"
    out.mkdir(parents=True, exist_ok=True)
    (out / "data_themes.js").write_text(
        '/* ---- Cyber-theme coverage, from the workbook\'s "Mapping cyber themes".\n'
        "   Generated by tools/cyber_themes.py - do not edit by hand.\n"
        "   Which themes each national framework covers, not how many rules it has.\n"
        "   Partial on purpose: the sheet holds %d of 29 countries. ---- */\n"
        "const CYBER_THEMES = %s;\n"
        % (len(countries), json.dumps(payload, ensure_ascii=False, indent=1)),
        encoding="utf-8")

    print("%d pays, %d thèmes répartis en %d familles"
          % (len(countries), len(themes), len(families)))
    print("  dernière mise à jour de la feuille : %s" % (payload["lastUpdate"] or "inconnue"))
    if unknown:
        print("  pays non reconnus : %s" % ", ".join(unknown))
    print()
    for iso, rec in sorted(countries.items(), key=lambda kv: -kv[1]["themesCovered"]):
        print("  %s %-16s %2d / %d thèmes couverts"
              % (iso, rec["country"], rec["themesCovered"], rec["themesTotal"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
