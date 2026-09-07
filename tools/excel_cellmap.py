#!/usr/bin/env python3
"""Build the address book of the comparative KPI workbook.

Every fact RegWatch shows lives in one cell of `CYBER WATCH5_Technical inventory.xlsx`,
addressed by (sheet, country row, field column). This script extracts that addressing
once, so two things become possible:

  1. the watch agent can name the exact cells a source would change, instead of
     emitting a vague "suggested action";
  2. RegWatch can read the workbook back and know which field it is looking at.

Output: data/excel-cellmap.json

    python3 tools/excel_cellmap.py "ressources/CYBER WATCH5_Technical inventory.xlsx"

Read-only: the workbook is never written back. Columns holding formulas are flagged
`writable: false` - those are computed by Excel and must never be overwritten.
"""

import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit("pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent

# The sheets that hold one row per country. Hidden helper sheets are excluded.
SHEETS = [
    "ID - P1",
    "Incident reporting - P1",
    "Registration - P1",
    "Cybersecurity frameworks - P1",
    "Audit & Controls - P1",
    "Sanctions - P2",
    "Authority - P3",
]

# Columns that are working notes, not comparative data.
SKIP_HEADERS = {"repartition", "répartition", "column1"}

# Workbook country label -> RegWatch ISO code.
ISO = {
    "austria": "AT", "belgium": "BE", "bulgaria": "BG", "croatia": "HR",
    "cyprus": "CY", "czechia": "CZ", "czech republic": "CZ", "denmark": "DK",
    "estonia": "EE", "finland": "FI", "france": "FR", "germany": "DE",
    "greece": "GR", "hungary": "HU", "ireland": "IE", "italy": "IT",
    "latvia": "LV", "lithuania": "LT", "luxembourg": "LU", "malta": "MT",
    "netherlands": "NL", "norway": "NO", "poland": "PL", "portugal": "PT",
    "romania": "RO", "slovakia": "SK", "slovenia": "SI", "spain": "ES",
    "sweden": "SE", "united kingdom": "GB", "uk": "GB",
}

# Aggregate rows at the bottom of each sheet - not countries, expected, not a warning.
AGGREGATES = {"total", "total/moyenne", "moyenne", "average"}


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def find_header_row(ws):
    """The header row is the one carrying the 'Country' label (row 3 or 4 here)."""
    for r in range(1, 12):
        for cell in ws[r]:
            if clean(cell.value).lower() == "country":
                return r, cell.column_letter
    return None, None


def build(path):
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)

    sheets = {}
    for name in SHEETS:
        if name not in values.sheetnames:
            print("  skipped (absent): %s" % name)
            continue
        ws, wf = values[name], formulas[name]
        hrow, ccol = find_header_row(ws)
        if not hrow:
            print("  skipped (no Country column): %s" % name)
            continue

        # country -> row
        rows = {}
        for r in range(hrow + 1, ws.max_row + 1):
            label = clean(ws["%s%d" % (ccol, r)].value)
            iso = ISO.get(label.lower())
            if iso:
                rows[iso] = r
            elif label and label.lower() not in AGGREGATES:
                print("  %s: unmapped country label %r (row %d)" % (name, label, r))

        # field -> column, with a writability flag taken from the first data row
        fields = []
        probe = min(rows.values()) if rows else hrow + 1
        for cell in ws[hrow]:
            header = clean(cell.value)
            if not header or cell.column_letter == ccol:
                continue
            if header.lower() in SKIP_HEADERS or re.fullmatch(r"\d+", header):
                continue
            raw = wf["%s%d" % (cell.column_letter, probe)].value
            fields.append({
                "column": cell.column_letter,
                "label": header,
                "writable": not (isinstance(raw, str) and raw.startswith("=")),
            })

        sheets[name] = {"headerRow": hrow, "countryColumn": ccol, "rows": rows, "fields": fields}
        print("  %-32s %2d countries x %2d fields" % (name, len(rows), len(fields)))

    return sheets


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    path = Path(sys.argv[1]).expanduser()
    if not path.exists():
        raise SystemExit("workbook not found: %s" % path)

    sheets = build(path)
    total = sum(len(s["rows"]) * len(s["fields"]) for s in sheets.values())
    writable = sum(len(s["rows"]) * sum(1 for f in s["fields"] if f["writable"])
                   for s in sheets.values())

    out = ROOT / "data" / "excel-cellmap.json"
    out.parent.mkdir(exist_ok=True)
    out.write_text(json.dumps({
        "workbook": path.name,
        "sheets": sheets,
        "addressableCells": total,
        "writableCells": writable,
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print("\n%d addressable cells (%d writable) -> %s"
          % (total, writable, out.relative_to(ROOT)))


if __name__ == "__main__":
    main()
